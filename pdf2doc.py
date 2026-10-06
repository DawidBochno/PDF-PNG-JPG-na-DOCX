"""PDF/PNG/JPG -> DOCX. Wsadowo, z zachowaniem ukladu i tabel.

PDF z tekstem (nie skan) idzie prosto do konwersji - najlepsza jakosc.
Obrazy i zeskanowane PDF-y przechodza przez OCR (Tesseract), z ktorego
program sam sklada DOCX (akapity, wciecia, jedna czcionka).

Okienko: wskaz pliki albo folder INPUT, wskaz OUTPUT, klikaj Konwertuj.
Opcje w okienku: jezyk OCR, zakres stron, pomijanie pustych stron.
Konsola: python pdf2doc.py INPUT_folder OUTPUT_folder
Test:    python pdf2doc.py --selftest
"""
import csv
import io
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import tkinter as tk
from collections import defaultdict
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import pymupdf
from pdf2docx import Converter

# Zbudowany .exe rozpakowuje zasoby do katalogu tymczasowego, ale foldery
# robocze maja lezec obok samego .exe - stad dwa osobne korzenie.
if getattr(sys, "frozen", False):
    BUNDLE = Path(sys._MEIPASS)
    APP_DIR = Path(sys.executable).parent
else:
    BUNDLE = APP_DIR = Path(__file__).parent

IN_DIR = APP_DIR / "INPUT"
OUT_DIR = APP_DIR / "OUTPUT"
TESSDATA_DIR = BUNDLE / "tessdata"
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}
DPI = 300
# Bez tego kazde wywolanie OCR mruga czarnym oknem konsoli (w .exe i pod pyw).
NO_WINDOW = subprocess.CREATE_NO_WINDOW if hasattr(subprocess, "CREATE_NO_WINDOW") else 0


def find_tesseract() -> str:
    """Szuka silnika OCR. Wersja dolaczona do paczki ma pierwszenstwo, potem
    systemowa - instalator bez praw administratora wrzuca ja do folderu
    uzytkownika, a nie do Program Files, stad kilka lokalizacji.
    """
    bundled = BUNDLE / "tesseract" / "tesseract.exe"
    if bundled.exists():
        return str(bundled)
    found = shutil.which("tesseract")
    if found:
        return found
    local = os.environ.get("LOCALAPPDATA", "")
    for cand in (
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
        Path(local) / "Programs" / "Tesseract-OCR" / "tesseract.exe",
        Path(local) / "Tesseract-OCR" / "tesseract.exe",
    ):
        if cand.exists():
            return str(cand)
    raise FileNotFoundError(
        "Brak Tesseracta (silnika OCR) - jest potrzebny do skanow i zdjec.\n"
        "Uruchom install.bat, albo zainstaluj recznie:\n"
        "https://github.com/UB-Mannheim/tesseract/wiki"
    )


# pewnosc OCR (0-100): slowa i linie ponizej to zwykle logo, pieczatka, podpis
MIN_WORD_CONF = 40
MIN_CONF = 50
BULLET = re.compile(r"^([-–•*]\s|\d{1,2}[.)]\s|[a-z][)]\s)")


# napis w okienku -> jezyki Tesseracta (modele w tessdata/)
LANGS = {
    "polski + angielski": "pol+eng",
    "tylko polski": "pol",
    "polski + angielski + niemiecki": "pol+eng+deu",
    "polski + angielski + ukrainski": "pol+eng+ukr",
}


def parse_pages(spec: str, count: int) -> list[int]:
    """'1-3, 5, 8-' -> numery stron od 0, tylko istniejace. Puste = wszystkie."""
    if not spec.strip():
        return list(range(count))
    pages = set()
    for part in spec.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        a, dash, b = part.partition("-")
        try:
            lo = int(a)
            hi = (int(b) if b.strip() else count) if dash else lo
        except ValueError:
            raise ValueError(f"zly zakres stron: '{part}' (przyklad: 1-3, 5)") from None
        if lo < 1 or hi < lo:
            raise ValueError(f"zly zakres stron: '{part}' (przyklad: 1-3, 5)")
        pages.update(range(lo - 1, min(hi, count)))
    if not pages:
        raise ValueError(f"plik ma {count} str. - brak stron z zakresu '{spec}'")
    return sorted(pages)


def _ocr_lines(png_bytes: bytes, scale: float, lang: str = "pol+eng") -> list[dict]:
    """OCR jednego obrazu -> linie tekstu {x0, x1, y0, y1, text} w punktach.

    Jezyk bierzemy z wlasnego folderu tessdata/ (przenosnie, bez uprawnien
    administratora) - stad TESSDATA_PREFIX zamiast --tessdata-dir, ktore w
    tesseract 5 psuje parsowanie nastepujacego po nim configfile.
    """
    env = {**os.environ, "TESSDATA_PREFIX": str(TESSDATA_DIR)}
    out = subprocess.run(
        [find_tesseract(), "-", "-", "-l", lang, "--dpi", str(DPI), "tsv"],
        input=png_bytes, check=True, capture_output=True, env=env,
        creationflags=NO_WINDOW,
    ).stdout.decode("utf8")
    words = defaultdict(list)
    for r in csv.DictReader(io.StringIO(out), delimiter="\t", quoting=csv.QUOTE_NONE):
        if r["level"] == "5" and r["text"].strip() and float(r["conf"]) >= MIN_WORD_CONF:
            words[(r["block_num"], r["par_num"], r["line_num"])].append(r)
    lines = []
    for ws in words.values():
        text = " ".join(w["text"] for w in ws)
        if sum(float(w["conf"]) for w in ws) / len(ws) < MIN_CONF or not any(c.isalnum() for c in text):
            continue
        lines.append({
            "x0": min(int(w["left"]) for w in ws) / scale,
            "x1": max(int(w["left"]) + int(w["width"]) for w in ws) / scale,
            "y0": min(int(w["top"]) for w in ws) / scale,
            "y1": max(int(w["top"]) + int(w["height"]) for w in ws) / scale,
            # mediana, bo pieczatka nachodzaca na linie zawyza jej obrys
            "h": sorted(int(w["height"]) for w in ws)[len(ws) // 2] / scale,
            "text": text,
        })
    return lines


def _text_lines(page) -> list[dict]:
    """Linie z cyfrowej strony (w PDF mieszanym) - w tym samym formacie co OCR."""
    lines = []
    for b in page.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            text = "".join(s["text"] for s in ln["spans"]).strip()
            if text:
                x0, y0, x1, y1 = ln["bbox"]
                lines.append({"x0": x0, "x1": x1, "y0": y0, "y1": y1, "h": y1 - y0, "text": text})
    return lines


def _paragraphs(lines: list[dict]) -> list[dict]:
    """Skleja linie w akapity wg geometrii: nowy akapit po wiekszej przerwie,
    po krotkiej linii (koniec akapitu), przy wcieciu albo punktorze."""
    if not lines:
        return []
    right = max(ln["x1"] for ln in lines)
    left = min(ln["x0"] for ln in lines)
    paras = []
    prev = None
    for ln in lines:  # kolejnosc czytania (OCR trzyma kolumny osobno)
        h = ln["h"]
        new = (
            prev is None
            or ln["y0"] <= prev["y0"]  # nastepna kolumna / blok obok
            or ln["y0"] - prev["y1"] > 0.8 * h
            or prev["x1"] < right - 0.15 * (right - left)
            or ln["x0"] > prev["x0"] + h
            or BULLET.match(ln["text"])
        )
        if new:
            paras.append({"x0": ln["x0"], "first_x0": ln["x0"], "x1": ln["x1"],
                          "gap": 0 if prev is None else ln["y0"] - prev["y0"],
                          "h": [h], "ys": [ln["y0"]], "text": ln["text"]})
        else:
            p = paras[-1]
            p["x0"] = min(p["x0"], ln["x0"])
            p["x1"] = max(p["x1"], ln["x1"])
            p["h"].append(h)
            p["ys"].append(ln["y0"])
            # przeniesienie wyrazu: "zdro-" + "wotnych"
            if p["text"].endswith("-") and ln["text"][:1].islower():
                p["text"] = p["text"][:-1] + ln["text"]
            else:
                p["text"] += " " + ln["text"]
        prev = ln
    for p in paras:
        p["left"] = left
    return paras


def ocr_to_docx(src: Path, docx_path: Path, lang: str = "pol+eng",
                pages_sel: list[int] | None = None, skip_blank: bool = True) -> None:
    """Skan/zdjecie -> DOCX z akapitami odtworzonymi z OCR (bez pdf2docx:
    z niewidocznego tekstu OCR robil losowe rozmiary czcionek i gubil bloki)."""
    from docx import Document
    from docx.shared import Pt

    doc = pymupdf.open(src)  # obraz otwiera sie juz obrocony wg EXIF
    page_w, page_h = doc[0].rect.width, doc[0].rect.height
    if src.suffix.lower() in IMAGE_EXTS:
        pix = pymupdf.Pixmap(str(src))  # rozdzielczosc oryginalu
        zoom = max(pix.width, pix.height) / max(page_w, page_h)
        # zdjecie nie ma rozmiaru w punktach - traktujemy je jak strone A4
        a4 = 595.3 / page_w
        unit, page_w, page_h = zoom / a4, 595.3, page_h * a4
    else:
        zoom = unit = DPI / 72
    pages = []
    for no in range(len(doc)) if pages_sel is None else pages_sel:
        page = doc[no]
        if page.get_text().strip():
            lines = _text_lines(page)
        else:
            png = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes("png")
            lines = _ocr_lines(png, unit, lang)
        if lines or not skip_blank:  # pusta strona = np. tyl skanu dwustronnego
            pages.append(_paragraphs(lines))
    doc.close()

    def median(xs, default):
        xs = sorted(xs)
        return xs[len(xs) // 2] if xs else default

    paras = [p for ps in pages for p in ps]
    body_h = median((h for p in paras for h in p["h"]), 12)
    # odstep miedzy liniami akapitu = 1.15 rozmiaru czcionki (Times, interlinia
    # pojedyncza); sama wysokosc linii OCR zawyza rozmiar przez akcenty i ogonki
    pitch = median((b - a for p in paras for a, b in zip(p["ys"], p["ys"][1:])), body_h)
    body_pt = min(max(round(pitch / 1.15 * 2) / 2, 8), 16)

    out = Document()
    out.styles["Normal"].font.name = "Times New Roman"
    out.styles["Normal"].font.size = Pt(body_pt)
    sec = out.sections[0]
    sec.page_width, sec.page_height = Pt(page_w), Pt(page_h)
    sec.top_margin = sec.bottom_margin = Pt(36)
    sec.left_margin = Pt(min(max(min((p["left"] for p in paras), default=72), 18), 108))
    sec.right_margin = Pt(min(max(page_w - max((p["x1"] for p in paras), default=0), 18), 108))
    for n, ps in enumerate(pages):
        if not ps:  # zachowana pusta strona
            out.add_paragraph().paragraph_format.page_break_before = bool(n)
        for i, p in enumerate(ps):
            par = out.add_paragraph()
            if i == 0 and n:
                par.paragraph_format.page_break_before = True
            fmt = par.paragraph_format
            fmt.space_after = Pt(0)
            fmt.space_before = Pt(min(max(p["gap"] - pitch, 0), 36))
            fmt.left_indent = Pt(max(p["x0"] - p["left"], 0))
            fmt.first_line_indent = Pt(p["first_x0"] - p["x0"])
            par.add_run(p["text"])
    if not pages:
        out.add_paragraph("(nie rozpoznano tekstu)")
    out.save(docx_path)


def convert_one(src: Path, out_dir: Path, lang: str = "pol+eng",
                pages: str = "", skip_blank: bool = True) -> Path:
    """pages: zakres jak w okienku ('1-3, 5'; puste = wszystkie). Obrazy maja
    jedna strone, wiec zakres ich nie dotyczy."""
    out_dir.mkdir(parents=True, exist_ok=True)
    docx = out_dir / f"{src.stem}.docx"
    if src.suffix.lower() in IMAGE_EXTS:
        ocr_to_docx(src, docx, lang, None, skip_blank)
        return docx
    with pymupdf.open(src) as doc:
        sel = parse_pages(pages, len(doc))
        # skan = strona bez tekstu, ale z obrazem (pusta bez obrazu to nie skan)
        scan = any(not doc[i].get_text().strip() and doc[i].get_images() for i in sel)
    if scan:
        ocr_to_docx(src, docx, lang, sel, skip_blank)
        return docx
    # cyfrowy PDF: pdf2docx zachowuje uklad i tabele - najlepsza jakosc
    c = Converter(str(src))
    try:
        c.convert(str(docx), pages=sel)
    finally:
        c.close()
    return docx


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PDF/PNG/JPG -> DOCX")
        self.geometry("720x580")
        self.files: list[Path] = []
        self.log_q: queue.Queue[str] = queue.Queue()

        IN_DIR.mkdir(exist_ok=True)
        self.out_var = tk.StringVar(value=str(OUT_DIR))

        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")

        ttk.Button(top, text="Wybierz pliki...", command=self.pick_files).grid(row=0, column=0, sticky="w")
        ttk.Button(top, text="Wybierz folder INPUT...", command=self.pick_folder).grid(row=0, column=1, padx=6)
        self.src_lbl = ttk.Label(top, text="nic nie wybrano", foreground="gray")
        self.src_lbl.grid(row=1, column=0, columnspan=3, sticky="w", pady=(6, 12))

        ttk.Label(top, text="Folder OUTPUT:").grid(row=2, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.out_var, width=62).grid(row=3, column=0, columnspan=2, sticky="we")
        ttk.Button(top, text="Zmien...", command=self.pick_out).grid(row=3, column=2, padx=6)

        opt = ttk.LabelFrame(self, text="Opcje", padding=8)
        opt.pack(fill="x", padx=10)
        self.lang_var = tk.StringVar(value=next(iter(LANGS)))
        self.pages_var = tk.StringVar()
        self.blank_var = tk.BooleanVar(value=True)
        ttk.Label(opt, text="Jezyk OCR:").grid(row=0, column=0, sticky="w")
        ttk.Combobox(opt, textvariable=self.lang_var, values=list(LANGS), state="readonly",
                     width=32).grid(row=0, column=1, sticky="w", padx=6)
        ttk.Label(opt, text="Strony:").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(opt, textvariable=self.pages_var, width=20).grid(row=1, column=1, sticky="w",
                                                                    padx=6, pady=(6, 0))
        ttk.Label(opt, text="np. 1-3, 5   (puste = wszystkie)", foreground="gray").grid(
            row=1, column=2, sticky="w", pady=(6, 0))
        ttk.Checkbutton(opt, text="Pomijaj puste strony (np. tyl skanu dwustronnego)",
                        variable=self.blank_var).grid(row=2, column=0, columnspan=3, sticky="w",
                                                      pady=(6, 0))

        self.btn = ttk.Button(self, text="Konwertuj", command=self.start)
        self.btn.pack(pady=10)
        self.bar = ttk.Progressbar(self, mode="determinate")
        self.bar.pack(fill="x", padx=10)
        self.log = tk.Text(self, height=14, wrap="none")
        self.log.pack(fill="both", expand=True, padx=10, pady=10)

        # domyslnie bierz to, co lezy w INPUT
        self.load(self._scan(IN_DIR), str(IN_DIR))
        self.after(100, self.drain)

    @staticmethod
    def _scan(folder: Path):
        exts = {".pdf", *IMAGE_EXTS}
        return sorted(p for p in folder.glob("*") if p.suffix.lower() in exts)

    def load(self, files, where):
        self.files = list(files)
        n = len(self.files)
        self.src_lbl.config(
            text=f"{n} plik(ow) z: {where}" if n else f"brak plikow PDF/PNG/JPG w: {where}",
            foreground="black" if n else "gray",
        )

    def pick_files(self):
        f = filedialog.askopenfilenames(
            title="Wybierz pliki",
            filetypes=[("PDF/PNG/JPG", "*.pdf *.png *.jpg *.jpeg"), ("Wszystkie", "*.*")],
        )
        if f:
            self.load([Path(x) for x in f], "wybor reczny")

    def pick_folder(self):
        d = filedialog.askdirectory(title="Folder ze zrodlowymi plikami")
        if d:
            self.load(self._scan(Path(d)), d)

    def pick_out(self):
        d = filedialog.askdirectory(title="Folder wyjsciowy")
        if d:
            self.out_var.set(d)

    def start(self):
        if not self.files:
            messagebox.showwarning("PDF -> DOCX", "Najpierw wskaz pliki albo folder.")
            return
        opts = {"lang": LANGS[self.lang_var.get()], "pages": self.pages_var.get(),
                "skip_blank": self.blank_var.get()}
        try:  # literowka w zakresie: komunikat od razu, nie przy kazdym pliku
            parse_pages(opts["pages"], 10**6)
        except ValueError as e:
            messagebox.showwarning("PDF -> DOCX", str(e))
            return
        self.btn.config(state="disabled")
        self.log.delete("1.0", "end")
        self.bar.config(value=0, maximum=len(self.files))
        threading.Thread(target=self.work, args=(list(self.files), Path(self.out_var.get()), opts),
                         daemon=True).start()

    def work(self, files, out_dir, opts):
        ok = 0
        for i, pdf in enumerate(files, 1):
            try:
                docx = convert_one(pdf, out_dir, **opts)
                ok += 1
                self.log_q.put(f"[{i}/{len(files)}] OK   {pdf.name} -> {docx.name}")
            except Exception as e:  # jeden zly plik nie moze zatrzymac reszty
                self.log_q.put(f"[{i}/{len(files)}] BLAD {pdf.name}: {e}")
            self.log_q.put(f"__progress__{i}")
        self.log_q.put(f"__done__Gotowe: {ok}/{len(files)} przekonwertowanych -> {out_dir}")

    def drain(self):
        while not self.log_q.empty():
            msg = self.log_q.get()
            if msg.startswith("__progress__"):
                self.bar.config(value=int(msg[12:]))
            elif msg.startswith("__done__"):
                self.log.insert("end", msg[8:] + "\n")
                self.btn.config(state="normal")
            else:
                self.log.insert("end", msg + "\n")
            self.log.see("end")
        self.after(100, self.drain)


def selftest():
    import tempfile
    import zipfile

    import pymupdf

    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / "t.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 100), "WZ nr: 1/05-2018")
        page.draw_rect(pymupdf.Rect(72, 150, 400, 180))
        page.draw_line(pymupdf.Point(236, 150), pymupdf.Point(236, 180))
        page.insert_text((80, 170), "produkt 1")
        page.insert_text((250, 170), "1.00")
        doc.save(src)
        doc.close()

        out = convert_one(src, td / "OUT")
        assert out.exists() and out.stat().st_size > 0, "brak pliku docx"
        xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf8")
        assert "1/05-2018" in xml, "zgubiony tekst"
        assert "<w:tbl>" in xml, "tabela nie odtworzona"

        # obraz -> OCR -> docx
        png = td / "t.png"
        img = pymupdf.open()
        p = img.new_page(width=400, height=150)
        p.insert_text((20, 60), "Zamowienie nr 42", fontsize=24)
        img[0].get_pixmap(dpi=200).save(png)
        img.close()

        out2 = convert_one(png, td / "OUT")
        assert out2.exists() and out2.stat().st_size > 0, "OCR: brak pliku docx"
        xml2 = zipfile.ZipFile(out2).read("word/document.xml").decode("utf8")
        assert "42" in xml2, f"OCR nie odczytal tekstu z obrazu: {xml2[:300]}"

        def text_page(doc, txt, rot=0):
            p = doc.new_page(width=400, height=150)
            p.insert_text((20, 60), txt, fontsize=24)
            return p.get_pixmap(matrix=pymupdf.Matrix(3, 3).prerotate(rot))

        # PDF mieszany: strona cyfrowa + skan -> tekst skanu nie moze zginac
        tmp = pymupdf.open()
        scan = text_page(tmp, "Faktura nr 777")
        mixed = pymupdf.open()
        mixed.new_page().insert_text((72, 100), "Pismo przewodnie")
        mixed.new_page().insert_image(pymupdf.Rect(0, 0, 400, 150), pixmap=scan)
        mixed.save(td / "mix.pdf")
        xml3 = zipfile.ZipFile(convert_one(td / "mix.pdf", td / "OUT")
                               ).read("word/document.xml").decode("utf8")
        assert "777" in xml3 and "przewodnie" in xml3, f"PDF mieszany: {xml3[:300]}"
        # skan: jedna czcionka dla calego tekstu (pdf2docx dawal kazdemu slowu inny rozmiar)
        assert "<w:sz " not in xml3, "skan: rozne rozmiary czcionek"

        # zdjecie z telefonu: zapisane obrocone + EXIF Orientation=6
        jpg = text_page(tmp, "Sygnatura 4321", rot=-90).tobytes("jpg")
        exif = (b"Exif\x00\x00MM\x00\x2a\x00\x00\x00\x08\x00\x01"
                b"\x01\x12\x00\x03\x00\x00\x00\x01\x00\x06\x00\x00\x00\x00\x00\x00")
        (td / "foto.jpg").write_bytes(
            jpg[:2] + b"\xff\xe1" + (len(exif) + 2).to_bytes(2, "big") + exif + jpg[2:])
        xml4 = zipfile.ZipFile(convert_one(td / "foto.jpg", td / "OUT")
                               ).read("word/document.xml").decode("utf8")
        assert "4321" in xml4, f"EXIF: zdjecie nieobrocone: {xml4[:300]}"

        # opcje: zakres stron
        assert parse_pages("", 3) == [0, 1, 2]
        assert parse_pages("1-2, 5; 2", 10) == [0, 1, 4]
        assert parse_pages("3-", 5) == [2, 3, 4]
        assert parse_pages("2-9", 3) == [1, 2]
        for bad in ("a", "0", "3-1", "1-x"):
            try:
                parse_pages(bad, 5)
                raise AssertionError(f"przepuszczony zly zakres: {bad}")
            except ValueError:
                pass
        xml5 = zipfile.ZipFile(convert_one(td / "mix.pdf", td / "OUT", pages="2")
                               ).read("word/document.xml").decode("utf8")
        assert "777" in xml5 and "przewodnie" not in xml5, f"zakres stron: {xml5[:300]}"

        # opcje: pusta strona skanu (bialy obraz) pomijana albo zachowana
        white = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 100), False)
        white.set_rect(white.irect, (255, 255, 255))
        duplex = pymupdf.open()
        duplex.new_page().insert_image(pymupdf.Rect(0, 0, 400, 150), pixmap=scan)
        duplex.new_page().insert_image(pymupdf.Rect(0, 0, 400, 150), pixmap=white)
        duplex.save(td / "duplex.pdf")
        for skip, breaks in ((True, 0), (False, 1)):
            x = zipfile.ZipFile(convert_one(td / "duplex.pdf", td / "OUT", skip_blank=skip)
                                ).read("word/document.xml").decode("utf8")
            assert x.count("<w:pageBreakBefore/>") == breaks, f"puste strony skip={skip}"

    # opcje: jezyki OCR - kazdy model musi byc w tessdata/
    for codes in LANGS.values():
        for code in codes.split("+"):
            assert (TESSDATA_DIR / f"{code}.traineddata").exists(), f"brak modelu {code}"

    # Okno musi dac sie zbudowac - w zbudowanym .exe brak bibliotek Tk objawia
    # sie inaczej niz w zwyklym Pythonie: program po prostu znika bez sladu.
    app = App()
    app.withdraw()
    app.update()
    assert app.title() == "PDF/PNG/JPG -> DOCX", f"zly tytul okna: {app.title()}"
    app.destroy()
    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


if __name__ == "__main__":
    args = sys.argv[1:]
    if args and args[0] == "--selftest":
        selftest()
    elif args:  # tryb konsolowy: INPUT [OUTPUT]
        src = Path(args[0])
        dst = Path(args[1]) if len(args) > 1 else OUT_DIR
        files = App._scan(src) if src.is_dir() else [src]
        for p in files:
            print(f"{p.name} -> {convert_one(p, dst)}")
    else:
        import aktualizacja
        app = App()
        aktualizacja.start(app, "DawidBochno/PDF-PNG-JPG-na-DOCX", "master", "pdf2doc.py")
        app.mainloop()
